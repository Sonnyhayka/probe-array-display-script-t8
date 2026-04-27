
####  Renew excel file with data values
#######################################
import xlsxwriter
import openpyxl
import time
from datetime import datetime
import string
import numpy as np
def create_excel_mace():
    try:
        macewbook=openpyxl.load_workbook('MACE_LOG.xlsx')
    except:
        ### Creating an excel file
        macewbook = xlsxwriter.Workbook('MACE_LOG.xlsx')
        macewsheet = macewbook.add_worksheet()

        ### Formatting excel file
        bold = macewbook.add_format({'bold':True})
        cell_format = macewbook.add_format()
        cell_format.set_align('right')
        commentlength=50
        macewsheet.set_column('A:A', 15)
        macewsheet.set_column('B:B', 14)
        macewsheet.set_column('D:D', commentlength)
        macewsheet.set_column('E:E', commentlength/2)
        macewsheet.set_column('F:F', commentlength/2)
        macewsheet.set_column('G:G', commentlength/2)
        macewsheet.set_column('H:H', commentlength/2)
        macewsheet.set_column('I:I', commentlength/2)
        macewsheet.set_column('J:J', commentlength/4)
        macewsheet.set_column('K:K', commentlength/4)
        macewsheet.set_column('L:L', commentlength/4)
        macewsheet.set_column('M:M', commentlength)
        macewsheet.set_column('N:N', commentlength)

        ### Creating titles
        macewsheet.write(0, 0, "DATE TIME", bold)
        macewsheet.write(0, 1, "APPLIES FOR SHOT NUMBER (Default is next shot)", bold)
        macewsheet.write(0, 2, "BASE PRESSURE (10^-6 Torr", bold)
        macewsheet.write(0, 3, "MACE COMMENT", bold)
        macewsheet.write(0, 4, "DIAGNOSTIC 1", bold)
        macewsheet.write(0, 5, "DIAGNOSTIC 2", bold)
        macewsheet.write(0, 6, "DIAGNOSTIC 3", bold)
        macewsheet.write(0, 7, "DIAGNOSTIC 4", bold)
        macewsheet.write(0, 8, "DIAGNOSTIC 5", bold)
        macewsheet.write(0, 9, "OTHER 1", bold)
        macewsheet.write(0, 10, "OTHER 2", bold)
        macewsheet.write(0, 11, "OTHER 3", bold)
        macewsheet.write(0, 12, "CONTROL CODE", bold)
        macewsheet.write(0, 13, "ANALYZE CODE", bold)     
        
def save_macelog_excel(macetext):
    create_excel_mace()
    macewbook=openpyxl.load_workbook('MACE_LOG.xlsx')
    macewsheet=macewbook['Sheet1']
            #writerow=macewsheet.max_row+1
            #macewsheet.cell(row=writerow, column=1).value=datetime.now().strftime('%m-%d-%Y_%H%M')
            #for i in range(1):#len(macetext)):
                #macewsheet.cell(row=writerow, column=i+2).value=macetext[i]
    row = str(macewsheet.max_row+1)
    col = list(string.ascii_uppercase)
    i = 0
    for key in macetext:
        if macetext[key] == 0 :
            macetext[key] = "--"
        if macetext[key] =='0':
            macetext[key] = "--"
        cell = col[i] + row
        macewsheet[cell] = macetext[key]
        i += 1 
    macewbook.save('MACE_LOG.xlsx')
    
def last_shot():
    import os
    try: 
        for root, dirs, files in os.walk("data\\", topdown=False):
            for i in np.arange(0, len(dirs)):
                row = int(dirs[i])
        lastshotname = "{:05d}".format(row+1)      # Takes the following shot number
 
    except:     #There are no files on folder
        lastshotname = "001"
    return lastshotname
